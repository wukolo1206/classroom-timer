# -*- coding: utf-8 -*-
"""全校版「📚 SR 閱讀小卡」分頁：本機驗收。

執行：python -m unittest discover -s tests -p test_sr_tab.py -v
只用虛構姓名；資料只在測試用的暫存瀏覽器 context。
"""
import functools, http.server, io, json, os, shutil, socketserver, tempfile, threading, unittest

import openpyxl
from playwright.sync_api import sync_playwright

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'universal.html')
NAMES = ['王小明', '李小華', '', '陳大同', '林美美', '張三']   # 第 3 行空白＝空號
CLASS = {'className': '408', 'teacherName': '示範老師', 'leaderName': '甲、乙', 'studentCount': 6, 'studentNames': NAMES}


def xlsx_bytes(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class SrTab(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix='ctu_sr_')
        os.makedirs(os.path.join(cls.tmp, 'classroom-timer'))
        shutil.copy(SRC, os.path.join(cls.tmp, 'classroom-timer', 'universal.html'))
        os.makedirs(os.path.join(cls.tmp, 'classroom-timer', 'sr-reading'))
        shutil.copy(os.path.join(ROOT, 'sr-reading', 'parent.html'), os.path.join(cls.tmp, 'classroom-timer', 'sr-reading', 'parent.html'))

        class Quiet(http.server.SimpleHTTPRequestHandler):
            def log_message(self, *a): pass
        cls.srv = socketserver.TCPServer(('127.0.0.1', 0), functools.partial(Quiet, directory=cls.tmp))
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()
        cls.url = 'http://127.0.0.1:%d/classroom-timer/universal.html' % cls.srv.server_address[1]
        cls.pw = sync_playwright().start()
        cls.browser = cls.pw.chromium.launch()

    @classmethod
    def tearDownClass(cls):
        cls.browser.close(); cls.pw.stop(); cls.srv.shutdown()
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def open(self, storage=None, cls_cfg=CLASS, seed_hidden=True):
        ctx = self.browser.new_context(accept_downloads=True, viewport={'width': 1366, 'height': 900})
        page = ctx.new_page()
        page._errs = []
        page.on('pageerror', lambda e: page._errs.append(str(e)))
        page._dialogs = []
        page.on('dialog', lambda d: (page._dialogs.append(d.message), d.accept()))
        init = {}
        if seed_hidden:
            init['ctu_timerTabHidden'] = '[]'
        if cls_cfg:
            init['ctu_class'] = json.dumps(cls_cfg, ensure_ascii=False)
        init['ctu_meta'] = json.dumps({'schema': 1, 'sh150Import': {'result': 'declined'}})
        init.update(storage or {})
        ctx.add_init_script('(() => { if (sessionStorage.getItem("__seeded")) return; sessionStorage.setItem("__seeded", "1");'
                            ' const s = %s; for (const k in s) localStorage.setItem(k, s[k]); })()' % json.dumps(init, ensure_ascii=False))
        page.goto(self.url)
        page.wait_for_timeout(600)
        self.addCleanup(ctx.close)
        return page

    def sr(self, page):
        page.click('#tab-sr-btn')
        page.wait_for_function("() => { const f = document.getElementById('sr-frame'); return f && f.contentWindow && f.contentWindow.renderAll && f.contentDocument.getElementById('roster-tbody').children.length > 0; }", timeout=8000)
        page.wait_for_timeout(300)
        return page.query_selector('#sr-frame').content_frame()

    def storage(self, page):
        return page.evaluate("() => { const o = {}; for (let i = 0; i < localStorage.length; i++) { const k = localStorage.key(i); o[k] = localStorage.getItem(k); } return o; }")

    def records(self, page):
        raw = self.storage(page).get('ctu_sr_records')
        return json.loads(raw) if raw else None

    def paste(self, fr, text):
        fr.evaluate("t => { document.getElementById('paste-textarea').value = t; parsePastedText(); }", text)

    # ── 分頁與名冊 ──
    def test_T01_tab_visible_by_default_and_for_existing_users(self):
        p = self.open(seed_hidden=False)
        self.assertTrue(p.is_visible('#tab-sr-btn'), '新老師預設要看得到 SR 分頁')
        p2 = self.open(storage={'ctu_timerTabHidden': '["timer","seat","group"]',
                                'ctu_timerTabOrder': '["sh150","weekly","timer","seat","group"]'})
        self.assertTrue(p2.is_visible('#tab-sr-btn'), '已存過分頁設定的老師也要看得到新分頁')

    def test_T02_roster_from_class_masked_and_skip_blank_seat(self):
        p = self.open()
        fr = self.sr(p)
        rows = fr.evaluate("[...document.querySelectorAll('#roster-tbody tr')].map(r => [r.children[1].innerText.trim(), r.children[2].innerText.trim()])")
        self.assertEqual([r[0] for r in rows], ['01', '02', '04', '05', '06'], '空行是空號，不列出')
        self.assertEqual(rows[0][1], '王○明')
        self.assertNotIn('王小明', fr.inner_text('#roster-tab'), '畫面上不可出現全名')
        self.assertEqual(p._errs, [])

    def test_T03_term_labels_follow_class_grade(self):
        p = self.open()
        fr = self.sr(p)
        opts = fr.evaluate("[...document.getElementById('term-select').options].map(o => o.text)")
        self.assertIn('115上（四上）（本學期）', opts)
        self.assertIn('115下（四下）', opts)

    # ── 貼上合併 ──
    def test_T04_paste_merges_by_seat_and_keeps_other_terms(self):
        p = self.open()
        fr = self.sr(p)
        self.paste(fr, '1 王小明 411-417\n2 李小華 2026/09/21 四年級 450–514\n3 某某 400\n9 路人 400\n5 林大美 380-440')
        rec = self.records(p)['bySeat']
        self.assertEqual(rec['1']['terms']['115-1']['sr'], '411-417')
        self.assertEqual(rec['2']['terms']['115-1']['date'], '2026/09/21')
        self.assertNotIn('3', rec, '空號不寫入')
        self.assertNotIn('9', rec, '名冊外座號不寫入')
        self.assertEqual(rec['5']['terms']['115-1']['sr'], '380-440')
        msg = p._dialogs[-1]
        self.assertIn('不在班級名冊', msg)
        self.assertIn('姓名不同', msg)
        self.assertNotIn('林大美', msg, '提示訊息的姓名也要遮罩')
        # 換到下學期貼上，不影響上學期
        fr.evaluate("switchTerm('115-2')")
        self.paste(fr, '1 王小明 440-480')
        rec = self.records(p)['bySeat']
        self.assertEqual(rec['1']['terms']['115-1']['sr'], '411-417')
        self.assertEqual(rec['1']['terms']['115-2']['sr'], '440-480')
        self.assertEqual(rec['1']['terms']['115-2']['gradeLevel'], '四年級')
        # 空白 SR 不覆蓋
        self.paste(fr, '1 王小明 -')
        self.assertEqual(self.records(p)['bySeat']['1']['terms']['115-2']['sr'], '440-480')

    def test_T05_edit_cell_keeps_focus_and_saves(self):
        p = self.open()
        fr = self.sr(p)
        inp = fr.locator('#roster-tbody input[data-row="0"][data-col="sr"]')
        inp.fill('394')
        inp.press('Enter')
        self.assertEqual(fr.evaluate("document.activeElement.dataset.row + '/' + document.activeElement.dataset.col"), '1/sr')
        self.assertIn('三年級', fr.inner_text('#cell-grade-0'))
        self.assertEqual(self.records(p)['bySeat']['1']['terms']['115-1']['sr'], '394')

    def test_T06_reload_keeps_data_and_term(self):
        p = self.open()
        fr = self.sr(p)
        fr.evaluate("switchTerm('115-2')")
        self.paste(fr, '2 李小華 420')
        p.reload(); p.wait_for_timeout(600)
        fr = self.sr(p)
        self.assertEqual(fr.evaluate("document.getElementById('term-select').value"), '115-2')
        self.assertEqual(fr.locator('#roster-tbody input[data-row="1"][data-col="sr"]').input_value(), '420')

    # ── 儲存隔離 ──
    def test_T07_only_ctu_keys_written(self):
        p = self.open(storage={'sh150_class_config': '{"className":"舊"}', 'sr408-roster-v1': '{"keep":1}'})
        before = self.storage(p)
        fr = self.sr(p)
        self.paste(fr, '1 王小明 411-417')
        after = self.storage(p)
        changed = [k for k in after if after.get(k) != before.get(k)]
        self.assertTrue(changed)
        self.assertTrue(all(k.startswith('ctu_') for k in changed), changed)
        self.assertEqual(after['sh150_class_config'], '{"className":"舊"}')
        self.assertEqual(after['sr408-roster-v1'], '{"keep":1}', '408 單機版的資料不可動')

    def test_T08_class_change_keeps_records_of_removed_seats(self):
        p = self.open()
        fr = self.sr(p)
        self.paste(fr, '6 張三 500')
        small = dict(CLASS, studentCount=5, studentNames=NAMES[:5])
        p.evaluate("c => { localStorage.setItem('ctu_class', JSON.stringify(c)); }", small)
        p.reload(); p.wait_for_timeout(600)
        fr = self.sr(p)
        self.assertEqual(self.records(p)['bySeat']['6']['terms']['115-1']['sr'], '500', '名冊變小不可刪紀錄')
        self.assertIn('不在目前名冊', fr.inner_text('#orphan-note'))

    # ── 匯出／匯入 ──
    def test_T09_export_then_import_roundtrip(self):
        p = self.open()
        fr = self.sr(p)
        self.paste(fr, '1 王小明 2026/09/21 四年級 411-417\n2 李小華 450-514')
        fr.evaluate("switchTerm('115-2')")
        self.paste(fr, '1 王小明 2027/03/02 四年級 440-480')
        before = self.records(p)['bySeat']
        with p.expect_download() as d:
            fr.evaluate("exportExcel()")
        path = os.path.join(self.tmp, 'export.xlsx')
        d.value.save_as(path)
        wb = openpyxl.load_workbook(path)
        ws = wb.worksheets[0]
        values = [c for row in ws.iter_rows(values_only=True) for c in row if c]
        self.assertIn('SR(115-1)', values)
        self.assertIn('王小明', values, '匯出檔給老師自己用，保留全名')
        # 清掉後再匯入
        p.evaluate("localStorage.removeItem('ctu_sr_records')")
        p.reload(); p.wait_for_timeout(600)
        fr = self.sr(p)
        fr.set_input_files('#excel-file-input', path)
        p.wait_for_timeout(800)
        after = self.records(p)['bySeat']
        for seat in ('1', '2'):
            for term, e in before[seat]['terms'].items():
                for f in ('sr', 'date'):
                    if e.get(f):
                        self.assertEqual(after[seat]['terms'][term].get(f), e[f], (seat, term, f))

    def test_T10_import_408_legacy_headers(self):
        p = self.open()
        fr = self.sr(p)
        data = xlsx_bytes([
            ['碧華國小 408 班 學生 SR 閱讀能力跨學期成長追蹤總表'],
            [],
            ['座號', '學生姓名', '測驗時間', '就讀年級', '四上SR', '適讀年級', '測驗時間', '就讀年級', '四下SR', '適讀年級'],
            ['01', '王小明', '2026/09/21', '四年級', '411-417', '四年級', '', '四年級', '', ''],
            ['02', '李小華', '2026/09/21', '四年級', '450-514', '', '2027/03/01', '四年級', '460-520', ''],
        ])
        fr.set_input_files('#excel-file-input', files=[{'name': 'old408.xlsx', 'mimeType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'buffer': data}])
        p.wait_for_timeout(800)
        rec = self.records(p)['bySeat']
        self.assertEqual(rec['1']['terms']['115-1']['sr'], '411-417')
        self.assertEqual(rec['2']['terms']['115-2']['sr'], '460-520')
        self.assertEqual(rec['2']['terms']['115-2']['date'], '2027/03/01')
        self.assertNotIn('115-2', rec['1']['terms'], '空白學期不建立紀錄')

    # ── 列印 ──
    def test_T11_print_cards_full_names_and_pages(self):
        p = self.open()
        fr = self.sr(p)
        self.paste(fr, '1 王小明 411-417\n2 李小華 450-514')
        fr.evaluate("selectOnlyTested()")
        fr.evaluate("showTab('print-tab')")
        txt = fr.inner_text('#print-section')
        self.assertIn('01 王小明', txt, '小卡印全名')
        self.assertIn('四～八年級', txt)
        self.assertEqual(fr.evaluate("document.querySelectorAll('.a4-sheet').length"), 1)
        fr.evaluate("document.getElementById('cfg-history-toggle').value='history'; renderPrintCards()")
        self.assertNotIn('需關懷', fr.inner_text('#print-section'))
        clipped = fr.evaluate("[...document.querySelectorAll('.sr-card')].filter(c => c.scrollHeight > c.clientHeight + 1).length")
        self.assertEqual(clipped, 0)
        # 框架內按列印時，瀏覽器只印框架自己的文件：把框架目前的畫面單獨拿出來印成 PDF 檢查
        fr.evaluate("document.querySelectorAll('.tab-content').forEach(el => el.classList.toggle('active', el.id === 'roster-tab'))")
        html = fr.evaluate("(() => { const d = document.documentElement.cloneNode(true); d.querySelectorAll('script').forEach(s => s.remove()); return '<!DOCTYPE html>' + d.outerHTML; })()")
        solo = p.context.new_page()
        solo.set_content(html)
        solo.emulate_media(media='print')
        self.assertGreater(solo.evaluate("document.getElementById('print-section').getBoundingClientRect().height"), 500,
                           '在名冊頁按列印也要印出小卡，不可空白')
        pdf = solo.pdf(prefer_css_page_size=True)
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(pdf))
        self.assertEqual(len(reader.pages), 1, '2 位學生＝1 張 A4，不多印白紙')
        # PDF 抽字會把「小」抽成部首字「⼩」，先正規化再比對
        import unicodedata
        self.assertIn('01 王小明', unicodedata.normalize('NFKC', reader.pages[0].extract_text()))

    def test_T12_parent_text_uses_class_settings(self):
        p = self.open()
        fr = self.sr(p)
        t = fr.evaluate("document.getElementById('parent-line-text').value")
        self.assertIn('408 導師 示範老師', t)
        self.assertNotIn('國榮', t)

    # ── 備份 ──
    def test_T13_full_backup_includes_sr(self):
        p = self.open()
        fr = self.sr(p)
        self.paste(fr, '1 王小明 411-417')
        with p.expect_download() as d:
            p.evaluate("ctuFullBackup('')")
        path = os.path.join(self.tmp, 'backup.json')
        d.value.save_as(path)
        with open(path, encoding='utf-8') as f:
            data = json.load(f)['data']
        self.assertIn('ctu_sr_records', data)

    def test_T14_no_names_lists_seats_by_count(self):
        p = self.open(cls_cfg={'className': '4XX', 'studentCount': 3, 'studentNames': []})
        fr = self.sr(p)
        rows = fr.evaluate("[...document.querySelectorAll('#roster-tbody tr')].map(r => r.children[1].innerText.trim())")
        self.assertEqual(rows, ['01', '02', '03'], '沒有姓名時依人數列座號')

    def test_T16_sample_files_import(self):
        roster = {'className': '408', 'teacherName': '示範老師', 'studentCount': 3, 'studentNames': ['王', '無', '李']}
        p = self.open(cls_cfg=roster)
        fr = self.sr(p)
        fr.set_input_files('#excel-file-input', os.path.join(ROOT, 'sr-reading', '範例', 'SR匯入範例_單一學期.xlsx'))
        p.wait_for_timeout(800)
        rec = self.records(p)['bySeat']
        self.assertEqual(rec['1']['terms']['115-1']['sr'], '411-417')
        self.assertEqual(rec['3']['terms']['115-1']['sr'], '394')
        self.assertNotIn('4', rec)
        self.assertIn('不在班級名冊', p._dialogs[-1])
        fr.set_input_files('#excel-file-input', os.path.join(ROOT, 'sr-reading', '範例', 'SR匯入範例_跨學期.xlsx'))
        p.wait_for_timeout(800)
        rec = self.records(p)['bySeat']
        self.assertEqual(rec['2']['terms']['115-2']['sr'], '470-530')
        self.assertEqual(rec['3']['terms']['115-1']['sr'], '311-398')
        self.assertNotIn('115-2', rec['3']['terms'])

    def test_T17_template_download_fill_import(self):
        p = self.open()
        fr = self.sr(p)
        with p.expect_download() as d:
            fr.evaluate("downloadTemplate()")
        path = os.path.join(self.tmp, 'template.xlsx')
        d.value.save_as(path)
        wb = openpyxl.load_workbook(path)
        ws = wb['SR登記']
        self.assertEqual([c.value for c in ws[4]], ['座號', '學生姓名', '測驗時間', '就讀年級', 'SR(115-1)'])
        self.assertEqual(ws['A5'].value, '01')
        self.assertEqual(ws['B5'].value, '王小明', '範本給老師填，帶全名')
        ws['E5'] = '411-417'
        ws['E6'] = '450-514'
        wb.save(path)
        fr.set_input_files('#excel-file-input', path)
        p.wait_for_timeout(800)
        rec = self.records(p)['bySeat']
        self.assertEqual(rec['1']['terms']['115-1']['sr'], '411-417')
        self.assertEqual(rec['2']['terms']['115-1']['sr'], '450-514')
        self.assertNotIn('姓名不同', p._dialogs[-1])

    def test_T18_wrong_roster_asks_before_writing(self):
        # 名冊還是 3 人測試名冊，卻匯入別班 5 人的檔案：先問，按取消就不寫
        roster = {'className': '408', 'teacherName': '示範老師', 'studentCount': 3, 'studentNames': ['王', '無', '李']}
        p = self.open(cls_cfg=roster)
        fr = self.sr(p)
        data = xlsx_bytes([['座號', '學生姓名', 'SR(115-1)'],
                           ['01', '甲同學', '400'], ['02', '乙同學', '410'], ['03', '丙同學', '420'],
                           ['04', '丁同學', '430'], ['05', '戊同學', '440']])
        fr.evaluate("() => { window.confirm = m => { window.__asked = m; return false; }; }")
        fr.set_input_files('#excel-file-input', files=[{'name': 'other.xlsx', 'mimeType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'buffer': data}])
        p.wait_for_timeout(800)
        asked = fr.evaluate("window.__asked || ''")
        self.assertIn('對不上', asked)
        self.assertIn('不寫入任何資料', asked)
        self.assertIsNone(self.records(p), '按取消不可寫入任何資料')
        # 名冊正確時不多問
        p2 = self.open()
        fr2 = self.sr(p2)
        fr2.evaluate("() => { window.__asked = ''; window.confirm = m => { window.__asked = m; return true; }; }")
        self.paste(fr2, '1 王小明 411-417\n2 李小華 450-514')
        self.assertEqual(fr2.evaluate("window.__asked"), '')
        self.assertEqual(self.records(p2)['bySeat']['2']['terms']['115-1']['sr'], '450-514')

    def test_T19_card_qr_default_change_and_clear(self):
        p = self.open()
        fr = self.sr(p)
        self.paste(fr, '1 王小明 411-417')
        fr.evaluate("showTab('print-tab')")
        self.assertTrue(fr.evaluate("!!document.querySelector('.sr-card .sr-card-qr svg')"), '預設要有 QR Code')
        self.assertTrue(fr.evaluate("document.getElementById('cfg-qr-url').value").endswith('/sr-reading/parent.html'))
        # 改網址會保存；重新整理後還在
        fr.fill('#cfg-qr-url', 'https://example.org/說明')
        p.reload(); p.wait_for_timeout(600)
        fr = self.sr(p)
        self.assertEqual(fr.evaluate("document.getElementById('cfg-qr-url').value"), 'https://example.org/說明')
        # 清空就不印 QR
        fr.evaluate("showTab('print-tab')")
        fr.fill('#cfg-qr-url', '')
        self.assertFalse(fr.evaluate("!!document.querySelector('.sr-card .sr-card-qr')"))
        fr.evaluate("resetQrUrl()")
        self.assertTrue(fr.evaluate("document.getElementById('cfg-qr-url').value").endswith('/sr-reading/parent.html'))

    def test_T20_parent_page_has_no_student_data_and_fits_phone(self):
        ctx = self.browser.new_context(viewport={'width': 375, 'height': 800})
        self.addCleanup(ctx.close)
        page = ctx.new_page()
        page.goto(self.url.replace('universal.html', 'sr-reading/parent.html'))
        self.assertFalse(page.evaluate("document.documentElement.scrollWidth > window.innerWidth"), '手機寬度不可左右捲動')
        self.assertIn('SR', page.inner_text('h1'))

    def test_T15_no_js_errors_across_tabs(self):
        p = self.open()
        fr = self.sr(p)
        for t in ('print-tab', 'parent-tab', 'norm-tab', 'roster-tab'):
            fr.evaluate("t => showTab(t)", t)
        fr.evaluate("document.getElementById('view-mode-select').value='comparison'; renderRosterTable()")
        p.click('#tab-weekly-btn'); p.click('#tab-sh150-btn'); p.click('#tab-sr-btn')
        p.wait_for_timeout(500)
        self.assertEqual(p._errs, [])


if __name__ == '__main__':
    unittest.main()
